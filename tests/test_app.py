from io import BytesIO
from pathlib import Path

import pandas as pd
from pandas.testing import assert_frame_equal
from streamlit.testing.v1 import AppTest

from app import filter_issues, to_csv_bytes, to_xlsx_bytes


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def test_filter_issues_combines_rule_column_and_text_filters():
    issues = pd.DataFrame({
        "Linha": [2, 3, 4],
        "Coluna": ["preco", "categoria", "preco"],
        "Código": ["R03", "R05", "R03"],
        "Valor encontrado": ["-10", "Eletronicos", "0"],
        "O que foi identificado": [
            "O preço deve ser maior que zero.",
            "A categoria não pertence à lista permitida.",
            "O preço deve ser maior que zero.",
        ],
    })

    result = filter_issues(issues, ["R03"], ["preco"], "-10")

    assert result["Linha"].tolist() == [2]


def test_to_csv_bytes_preserves_columns_and_content():
    dataframe = pd.DataFrame({"Regra": ["R03"], "Coluna": ["preco"]})

    result = pd.read_csv(BytesIO(to_csv_bytes(dataframe)))

    assert result.to_dict("records") == [{"Regra": "R03", "Coluna": "preco"}]


def test_streamlit_app_renders_empty_state_without_exceptions():
    app = AppTest.from_file(PROJECT_ROOT / "app.py")
    app.run()

    assert not app.exception
    assert len(app.sidebar.file_uploader) == 1


def test_to_xlsx_bytes_creates_readable_workbook():
    from openpyxl import load_workbook

    content = to_xlsx_bytes(pd.DataFrame({"Produto": ["Teclado"], "Preço": [12.5]}))
    workbook = load_workbook(BytesIO(content), read_only=True)

    assert workbook.sheetnames == ["Dados"]
    assert list(workbook.active.values)[0] == ("Produto", "Preço")
    assert list(workbook.active.values)[1] == ("Teclado", 12.5)


def test_exports_do_not_mutate_the_original_dataframe():
    source = pd.DataFrame({"Produto": ["Teclado"], "Preço": [12.5]})
    original_snapshot = source.copy(deep=True)

    to_csv_bytes(source)
    to_xlsx_bytes(source, "Produtos")

    assert_frame_equal(source, original_snapshot)


def test_streamlit_analysis_shows_expected_count_and_all_export_formats():
    def sample_analysis():
        from pathlib import Path

        from app import render_analysis
        from data_quality_agent.file_reader import load_file

        dataframe = load_file(Path("data/samples/products_with_issues.csv"))
        render_analysis(dataframe, "products_with_issues.csv")

    app = AppTest.from_function(sample_analysis).run(timeout=10)

    assert not app.exception
    assert app.metric[1].value == "8"
    assert len(app.download_button) == 10
